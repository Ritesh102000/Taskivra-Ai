"""Validate generated fixture-only exports with independent installed parsers."""
import hashlib
import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import fitz
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

root = Path(sys.argv[1]).resolve()
proof = json.loads((root / 'electron-proof.json').read_text())
for name, expected in proof['files'].items():
    data = (root / name).read_bytes()
    assert len(data) == expected['bytes']
    assert hashlib.sha256(data).hexdigest() == expected['sha256']

archive_parts = {}
for name in ['report.docx', 'report.xlsx']:
    with zipfile.ZipFile(root / name) as archive:
        assert archive.testzip() is None
        archive_parts[name] = archive.namelist()
        for part in archive.namelist():
            ElementTree.fromstring(archive.read(part))
            assert b'TargetMode="External"' not in archive.read(part)

report = PdfReader(root / 'report.pdf')
text = '\n'.join(page.extract_text() for page in report.pages)
for expected in ['Quarterly operations brief', 'North', 'South', '15,000', '+25%', 'Sources and coverage', 'Limitations', proof['sourceVersion'], proof['sourceSha256']]:
    assert expected in text, repr(expected)
probe_text = '\n'.join(page.extract_text() for page in PdfReader(root / 'security-probe.pdf').pages)
assert 'JAVASCRIPT_DISABLED_SENTINEL' in probe_text
assert 'JAVASCRIPT_EXECUTED_SENTINEL' not in probe_text
assert 'NETWORK WAS CONTACTED' not in probe_text
assert proof['networkRequests'] == 0

document = Document(root / 'report.docx')
paragraphs = '\n'.join(paragraph.text for paragraph in document.paragraphs)
for expected in ['Quarterly operations brief', 'two regions', 'Sources and coverage', 'Limitations', proof['sourceVersion'], proof['sourceSha256']]:
    assert expected in paragraphs, repr(expected)
assert len(document.tables) == 1
assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [
    ['Region', 'Previous', 'Current', 'Change'], ['North', '12,000', '15,000', '+25%'], ['South', '8,000', '8,800', '+10%']]

workbook = load_workbook(root / 'report.xlsx', data_only=False)
assert workbook.sheetnames == ['Results', 'Source record']
rows = list(workbook['Results'].values)
assert rows == [('Region', 'Previous', 'Current', 'Identifier', 'Literal formula'), ('North', '12000', '15000', '00123', '=SUM(B2:C2)'), ('South', '8000', '8800', '00045', '+SUM(B3:C3)')]
assert all(cell.data_type == 's' for row in workbook['Results'] for cell in row)
assert workbook['Results'].freeze_panes == 'A2'
assert workbook['Results'].auto_filter.ref == 'A1:E3'
assert proof['csvSourceSha256'] in str(list(workbook['Source record'].values))
assert not workbook._external_links

images = []
for name in ['report.pdf', 'word-render/report.pdf', 'sheet-render/report.pdf']:
    path = root / name
    if not path.exists():
        continue
    with fitz.open(path) as rendered:
        for index, page in enumerate(rendered):
            image = root / (name.replace('/', '-').replace('.pdf', '') + f'-page-{index + 1}.png')
            page.get_pixmap(matrix=fitz.Matrix(1.3, 1.3), alpha=False).save(image)
            images.append(str(image.relative_to(root)))

result = {'fixtureOnly': True, 'validated': True, 'pdfPages': len(report.pages), 'wordParagraphs': len(document.paragraphs), 'wordTables': len(document.tables), 'spreadsheetRows': len(rows) - 1, 'spreadsheetLiteralCells': 15, 'networkRequests': proof['networkRequests'], 'javascriptExecutionObserved': False, 'archiveParts': archive_parts, 'renderedImages': images}
(root / 'parser-proof.json').write_text(json.dumps(result, indent=2) + '\n')
print(json.dumps(result, indent=2))
